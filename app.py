import os
import json
from io import BytesIO

import streamlit as st
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document


# ==========================================
# 1. PAGE CONFIGURATION
# ==========================================

st.set_page_config(
    page_title="AI Resume Analyzer",
    page_icon="📄",
    layout="wide"
)

load_dotenv()

api_key = (
    os.getenv("GROQ_API_KEY")
    or os.getenv("GROQ_API_KEY")
)


# ==========================================
# 2. AI CONFIGURATION
# ==========================================

MODEL = "openai/gpt-oss-20b"


# ==========================================
# 3. DATA MODELS
# ==========================================

class JobD(BaseModel):
    role: str
    required_skills: list[str] = Field(default_factory=list)
    preferred_skills: list[str] = Field(default_factory=list)
    minimum_experience: float | None = None
    education_requirements: list[str] = Field(
        default_factory=list
    )
    responsibilities: list[str] = Field(
        default_factory=list
    )


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = Field(default_factory=list)


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = Field(default_factory=list)
    experiences: list[Experience] = Field(
        default_factory=list
    )
    education: list[str] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(
        default_factory=list
    )


class MatchResult(BaseModel):
    score: float
    details: dict


# ==========================================
# 4. AI CLIENT
# ==========================================

@st.cache_resource
def get_client(key):
    return Groq(api_key=key)


def ask_ai(client, system_prompt, user_prompt):

    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_prompt
            }
        ],
        response_format={
            "type": "json_object"
        },
        temperature=0.1
    )

    content = response.choices[0].message.content

    return json.loads(content)


# ==========================================
# 5. READ PDF RESUME
# ==========================================

def read_pdf(file_bytes):

    reader = PdfReader(BytesIO(file_bytes))

    text = ""

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


# ==========================================
# 6. READ WORD RESUME
# ==========================================

def read_docx(file_bytes):

    document = Document(BytesIO(file_bytes))

    text = ""

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            row_text = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if row_text:
                text += " | ".join(row_text) + "\n"

    return text


# ==========================================
# 7. EXTRACT JOB DESCRIPTION
# ==========================================

def analyze_job(client, job_description):

    system_prompt = """
    You are an expert HR assistant.

    Analyze the job description and extract the following fields:

    1. role: string
    2. required_skills: list of strings
    3. preferred_skills: list of strings
    4. minimum_experience: number or null
    5. education_requirements: list of strings
    6. responsibilities: list of strings

    Return ONLY a JSON object containing actual extracted data.

    IMPORTANT:
    - Do not return a JSON schema.
    - Do not return properties, title, type, or items fields.
    - Every field must contain actual extracted information.
    - role must be a string, not an object.
    - Skills, education requirements, and responsibilities
      must be lists of strings.
    - minimum_experience must be a number or null.
    - If minimum experience is not mentioned, return null.
    - Do not invent information.

    Example output:

    {
        "role": "Software Development Engineer I",
        "required_skills": ["Python", "Data Structures"],
        "preferred_skills": ["AWS", "SQL"],
        "minimum_experience": null,
        "education_requirements": [
            "Bachelor's degree in Computer Science"
        ],
        "responsibilities": [
            "Develop and maintain software"
        ]
    }
    """

    data = ask_ai(
        client,
        system_prompt,
        job_description
    )

    return JobD.model_validate(data)

# ==========================================
# 8. PARSE RESUME
# ==========================================

def parse_resume(client, resume_text):

    schema = Resume.model_json_schema()

    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume.

    Return only valid JSON matching this schema:

    {json.dumps(schema)}

    Rules:
    - Extract the candidate's name, email and phone.
    - Extract skills from the entire resume.
    - Include internships under experiences.
    - Extract education, projects and certifications.
    - Extract total experience in years if available.
    - Do not invent information.
    - Use null for unavailable single values.
    - Use empty lists for missing list values.
    """

    data = ask_ai(
        client,
        system_prompt,
        resume_text
    )

    return Resume.model_validate(data)


# ==========================================
# 9. MATCH RESUME WITH JOB
# ==========================================

def final_score(client, job, resume):

    system_prompt = """
    You are an expert technical recruiter.

    Compare the candidate's resume with the job description.

    Return valid JSON with these fields:

    {
        "score": 0,
        "details": {
            "candidate_name": "",
            "matching_skills": [],
            "missing_important_skills": [],
            "experience_requirement_met": "",
            "education_match": "",
            "relevant_projects": [],
            "final_verdict": ""
        }
    }

    Rules:
    - Score must be between 0 and 100.
    - Consider required skills, experience, education
      and relevant projects.
    - Identify skills that are present in the resume.
    - Identify important skills missing from the resume.
    - Do not invent qualifications or experience.
    - Do not treat preferred skills as mandatory.
    - If information is unavailable, say so.
    - Give a concise final verdict.
    """

    user_prompt = f"""
    JOB DESCRIPTION ANALYSIS:

    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:

    {resume.model_dump_json(indent=2)}

    Compare the candidate with this job.
    Return the matching analysis as JSON.
    """

    data = ask_ai(
        client,
        system_prompt,
        user_prompt
    )

    result = MatchResult.model_validate(data)

    # Keep the score within the valid range.
    result.score = max(0, min(100, result.score))

    return result


# ==========================================
# 10. WEBSITE DESIGN
# ==========================================

st.markdown(
    """
    <style>
    /* Overall page */
    .stApp {
        background:
            radial-gradient(ellipse at 8% 0%, rgba(124, 58, 237, 0.18), transparent 34%),
            radial-gradient(ellipse at 95% 12%, rgba(14, 165, 233, 0.10), transparent 30%),
            #0b1020;
        color: #f4f4fb;
    }

    [data-testid="stHeader"] {
        background: rgba(11, 16, 32, 0.78);
    }

    .block-container {
        max-width: 1180px;
        padding-top: 2rem;
        padding-bottom: 3rem;
    }

    /* Hero banner */
    .hero {
        padding: 2.3rem 2rem 2rem;
        margin: 0 0 1.8rem;
        border: 1px solid rgba(167, 139, 250, 0.28);
        border-radius: 26px;
        background: linear-gradient(125deg, rgba(109, 40, 217, 0.24), rgba(30, 41, 75, 0.72) 55%, rgba(14, 165, 233, 0.12));
        box-shadow: 0 18px 55px rgba(0, 0, 0, 0.22);
        text-align: center;
    }

    .hero-badge {
        display: inline-block;
        padding: 0.38rem 0.85rem;
        margin-bottom: 0.9rem;
        border: 1px solid rgba(196, 181, 253, 0.35);
        border-radius: 999px;
        background: rgba(124, 58, 237, 0.17);
        color: #ddd6fe;
        font-size: 0.82rem;
        font-weight: 650;
        letter-spacing: 0.04em;
    }

    .main-title {
        margin: 0;
        font-size: clamp(2rem, 5vw, 3.25rem);
        line-height: 1.15;
        font-weight: 850;
        letter-spacing: -0.045em;
        color: #ffffff;
    }

    .main-title span {
        background: linear-gradient(90deg, #c4b5fd, #a5f3fc);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
    }

    .subtitle {
        max-width: 700px;
        margin: 0.9rem auto 0;
        color: #c2c8dc;
        font-size: 1.04rem;
        line-height: 1.7;
    }

    .section-title {
        color: #e4ddff;
        font-size: 1.25rem;
        font-weight: 750;
        margin-bottom: 0.5rem;
    }

    .eyebrow {
        color: #a5b4fc;
        font-size: 0.78rem;
        font-weight: 750;
        letter-spacing: 0.12em;
        text-transform: uppercase;
        margin-bottom: 0.35rem;
    }

    /* Sidebar */
    [data-testid="stSidebar"] {
        background: linear-gradient(180deg, #11162a 0%, #0c1120 100%);
        border-right: 1px solid rgba(167, 139, 250, 0.15);
    }

    [data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p {
        color: #c4c9dc;
    }

    /* Input components */
    [data-testid="stFileUploader"] section {
        background: rgba(20, 27, 49, 0.82);
        border: 1px dashed #7164b5;
        border-radius: 18px;
        padding: 1rem;
    }

    [data-testid="stFileUploader"] section:hover {
        border-color: #a78bfa;
    }

    [data-testid="stTextArea"] textarea,
    [data-testid="stTextInput"] input {
        background: #11182b;
        color: #f4f4fb;
        border: 1px solid #343c59;
        border-radius: 14px;
    }

    [data-testid="stTextArea"] textarea:focus,
    [data-testid="stTextInput"] input:focus {
        border-color: #a78bfa;
        box-shadow: 0 0 0 1px #a78bfa;
    }

    /* Buttons */
    .stButton > button,
    .stDownloadButton > button {
        border: 1px solid rgba(196, 181, 253, 0.35);
        border-radius: 13px;
        min-height: 3rem;
        font-weight: 750;
        color: #ffffff;
        background: linear-gradient(95deg, #6d28d9, #7c3aed 55%, #4f46e5);
        box-shadow: 0 8px 24px rgba(109, 40, 217, 0.20);
        transition: all 0.2s ease;
    }

    .stButton > button:hover,
    .stDownloadButton > button:hover {
        border-color: #c4b5fd;
        transform: translateY(-1px);
        box-shadow: 0 12px 28px rgba(109, 40, 217, 0.32);
        color: #ffffff;
    }

    /* Metrics and result surfaces */
    [data-testid="stMetric"] {
        background: linear-gradient(145deg, rgba(26, 34, 59, 0.95), rgba(18, 24, 43, 0.95));
        border: 1px solid rgba(129, 140, 248, 0.22);
        border-radius: 18px;
        padding: 1rem 1.15rem;
    }

    [data-testid="stMetricLabel"] {
        color: #b8c0d9;
    }

    [data-testid="stMetricValue"] {
        color: #ddd6fe;
        font-weight: 800;
    }

    [data-testid="stProgressBar"] > div > div {
        background: linear-gradient(90deg, #7c3aed, #22d3ee);
    }

    .result-box {
        background: linear-gradient(135deg, rgba(30, 41, 75, 0.9), rgba(30, 27, 75, 0.75));
        padding: 1.25rem 1.4rem;
        border-radius: 18px;
        border: 1px solid rgba(167, 139, 250, 0.28);
        margin-bottom: 1rem;
    }

    [data-testid="stExpander"] {
        border: 1px solid rgba(129, 140, 248, 0.22);
        border-radius: 15px;
        background: rgba(17, 24, 43, 0.65);
    }

    hr {
        border-color: rgba(148, 163, 184, 0.16);
    }

    h1, h2, h3 {
        letter-spacing: -0.025em;
    }

    @media (max-width: 700px) {
        .block-container {
            padding-top: 1rem;
        }
        .hero {
            padding: 1.5rem 1rem;
            border-radius: 20px;
        }
        .subtitle {
            font-size: 0.95rem;
        }
    }
    </style>
    """,
    unsafe_allow_html=True
)


# ==========================================
# 11. SIDEBAR
# ==========================================

with st.sidebar:

    st.markdown("# ✦ CareerLens")
    st.caption("AI RESUME INTELLIGENCE")
    st.write("Turn your resume into a clearer career strategy.")

    st.divider()

    st.markdown("### YOUR WORKSPACE")
    st.write("📄  Resume screening")
    st.write("🧠  AI-powered insights")
    st.write("🎯  Job compatibility")
    st.write("🧩  Skill gap analysis")
    st.write("📥  Downloadable report")

    st.divider()
    st.info("Tip: Use the full job description for a more useful comparison.")
    st.caption("Built with Python · Streamlit · Groq")


# ==========================================
# 12. MAIN WEBSITE
# ==========================================

st.markdown(
    """
    <div class="hero">
        <div class="hero-badge">✦ AI-POWERED CAREER TOOL</div>
        <div class="main-title">Your next role starts with a <span>stronger resume.</span></div>
        <div class="subtitle">
            Compare your resume with a job description, uncover your strongest matches,
            and discover the skills you can improve — all in one place.
        </div>
    </div>
    """,
    unsafe_allow_html=True
)

st.markdown(
    '<div class="eyebrow">01 &nbsp; / &nbsp; Start your analysis</div>',
    unsafe_allow_html=True
)

left, right = st.columns(2)

with left:

    st.markdown(
        '<div class="section-title">📄 Upload your resume</div>',
        unsafe_allow_html=True
    )
    st.caption("PDF or DOCX · Keep your resume readable")

    uploaded_file = st.file_uploader(
        "Choose your resume",
        type=["pdf", "docx"],
        help="Upload a PDF or Word document."
    )

    if uploaded_file:
        st.success(
            f"Uploaded: {uploaded_file.name}"
        )


with right:

    st.markdown(
        '<div class="section-title">💼 Target job description</div>',
        unsafe_allow_html=True
    )
    st.caption("Paste the role requirements and responsibilities")

    job_description = st.text_area(
        "Paste the job description here",
        height=220,
        placeholder=(
            "Paste the complete job description here..."
        )
    )


st.write("")
st.caption("🔒 Your uploaded file is analyzed for this session; avoid uploading documents you do not have permission to share.")

analyze_button = st.button(
    "🚀 Analyze My Resume",
    use_container_width=True,
    type="primary"
)


# ==========================================
# 13. ANALYSIS PROCESS
# ==========================================

if analyze_button:

    if not api_key:

        st.error(
            "Groq API key not found. "
            "Please add GROQ_API_KEY to your .env file."
        )

    elif not uploaded_file:

        st.warning(
            "Please upload your resume first."
        )

    elif not job_description.strip():

        st.warning(
            "Please enter a job description."
        )

    else:

        client = get_client(api_key)

        try:

            with st.spinner("Reading your resume..."):

                file_bytes = uploaded_file.getvalue()

                if uploaded_file.name.lower().endswith(".pdf"):
                    resume_text = read_pdf(file_bytes)

                else:
                    resume_text = read_docx(file_bytes)

                if not resume_text.strip():

                    st.error(
                        "No readable text found in the resume. "
                        "If it is a scanned PDF, OCR may be needed."
                    )

                    st.stop()

            with st.spinner(
                "AI is analyzing the job description..."
            ):

                job = analyze_job(
                    client,
                    job_description
                )

            with st.spinner(
                "AI is extracting your resume details..."
            ):

                resume = parse_resume(
                    client,
                    resume_text
                )

            with st.spinner(
                "Calculating your job match score..."
            ):

                result = final_score(
                    client,
                    job,
                    resume
                )

            st.success("Resume analysis completed!")

            st.divider()

            # ----------------------------------
            # SCORE
            # ----------------------------------

            st.markdown("## 📊 Analysis Results")

            col1, col2, col3 = st.columns(3)

            with col1:

                st.metric(
                    "Resume Match",
                    f"{result.score:.1f}%"
                )

            with col2:

                st.metric(
                    "Required Skills",
                    len(job.required_skills)
                )

            with col3:

                st.metric(
                    "Your Skills",
                    len(resume.skills)
                )

            st.progress(
                int(result.score) / 100
            )

            details = result.details

            # ----------------------------------
            # CANDIDATE DETAILS
            # ----------------------------------

            st.markdown("## 👤 Candidate Profile")

            profile_col1, profile_col2 = st.columns(2)

            with profile_col1:

                st.write(
                    "**Name:**",
                    resume.name or "Not found"
                )

                st.write(
                    "**Email:**",
                    resume.email or "Not found"
                )

                st.write(
                    "**Phone:**",
                    resume.phone or "Not found"
                )

            with profile_col2:

                experience = resume.total_experience_years

                st.write(
                    "**Experience:**",
                    (
                        f"{experience} years"
                        if experience is not None
                        else "Not specified"
                    )
                )

                st.write(
                    "**Job Role:**",
                    job.role
                )

            # ----------------------------------
            # SKILLS
            # ----------------------------------

            st.markdown("## ✅ Matching Skills")

            matching_skills = details.get(
                "matching_skills",
                []
            )

            if matching_skills:

                st.write(
                    ", ".join(map(str, matching_skills))
                )

            else:

                st.info(
                    "No matching skills were identified."
                )

            st.markdown("## ❌ Skills to Improve")

            missing_skills = details.get(
                "missing_important_skills",
                []
            )

            if missing_skills:

                for skill in missing_skills:

                    st.write(f"• {skill}")

            else:

                st.success(
                    "No missing important skills were identified."
                )

            # ----------------------------------
            # EXPERIENCE AND EDUCATION
            # ----------------------------------

            st.markdown("## 🎓 Education & Experience")

            st.write(
                "**Education Match:**",
                details.get(
                    "education_match",
                    "Not available"
                )
            )

            st.write(
                "**Experience Requirement:**",
                details.get(
                    "experience_requirement_met",
                    "Not available"
                )
            )

            st.markdown("## 💡 Relevant Projects")

            relevant_projects = details.get(
                "relevant_projects",
                []
            )

            if relevant_projects:

                for project in relevant_projects:

                    st.write(f"• {project}")

            else:

                st.write(
                    "No relevant projects identified by the AI."
                )

            # ----------------------------------
            # COMPLETE RESUME
            # ----------------------------------

            with st.expander(
                "View Extracted Resume Details"
            ):

                st.markdown("### Skills")

                st.write(
                    resume.skills or ["No skills extracted"]
                )

                st.markdown("### Education")

                st.write(
                    resume.education or ["No education extracted"]
                )

                st.markdown("### Projects")

                st.write(
                    resume.projects or ["No projects extracted"]
                )

                st.markdown("### Certifications")

                st.write(
                    resume.certifications
                    or ["No certifications extracted"]
                )

                st.markdown("### Work Experience")

                if resume.experiences:

                    for exp in resume.experiences:

                        st.markdown(
                            f"**{exp.role or 'Role not specified'}**"
                        )

                        st.write(
                            exp.company or "Company not specified"
                        )

                        st.write(
                            exp.description or "No description available"
                        )

                        st.write(
                            "Skills:",
                            exp.skills_used
                        )

                        st.divider()

                else:

                    st.write(
                        "No work experience extracted."
                    )

            # ----------------------------------
            # FINAL VERDICT
            # ----------------------------------

            st.markdown("## 🎯 AI Verdict")

            st.markdown(
                '<div class="result-box">',
                unsafe_allow_html=True
            )

            st.write(
                details.get(
                    "final_verdict",
                    "No final verdict available."
                )
            )

            st.markdown(
                "</div>",
                unsafe_allow_html=True
            )

            # ----------------------------------
            # DOWNLOAD REPORT
            # ----------------------------------

            report = {
                "candidate": resume.model_dump(),
                "job": job.model_dump(),
                "match_score": result.score,
                "analysis": result.details
            }

            st.download_button(
                label="⬇️ Download Analysis Report",
                data=json.dumps(
                    report,
                    indent=2,
                    ensure_ascii=False
                ),
                file_name="resume_analysis.json",
                mime="application/json",
                use_container_width=True
            )

            st.caption(
                "AI-generated results may contain errors. "
                "Verify important qualifications manually."
            )

        except Exception as error:

            st.error(
                "Something went wrong during analysis."
            )

            st.exception(error)


# ==========================================
# 14. FOOTER
# ==========================================

st.divider()

st.markdown(
    """
    <div style="text-align:center; color:#9292a5;">
        AI Resume Analyzer | Built with Python and Streamlit
    </div>
    """,
    unsafe_allow_html=True
)
